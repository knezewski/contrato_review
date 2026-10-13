"""Text extraction from PDF and DOCX (data-model.md §1).

Failure modes (constitution IV):
- The file type is determined by signature (`%PDF`, a DOCX ZIP container), not by extension.
- A PDF without pages or with a broken structure → DocumentUnreadable.
- A PDF page whose text layer is empty after strip() marks the document as needs_ocr.
  Scanned text is not recognized here: the caller runs OCR if it is enabled.
- A DOCX that python-docx cannot open → DocumentUnreadable.
The file is read from memory (bytes) and never written to disk.
"""

from dataclasses import dataclass
from io import BytesIO

import docx
import pymupdf

from src.ingest.errors import DocumentUnreadable, UnsupportedFormat


@dataclass
class ContractDocument:
    source_format: str  # "pdf" | "docx"
    text_es: str
    page_count: int
    needs_ocr: bool
    ocr_applied: bool = False


def _extract_pdf(data: bytes) -> ContractDocument:
    try:
        pdf = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:  # pymupdf raises different exceptions on broken files
        raise DocumentUnreadable("PDF cannot be opened") from exc
    try:
        if pdf.page_count == 0:
            raise DocumentUnreadable("PDF without pages")
        pages = [page.get_text() for page in pdf]
        needs_ocr = any(not text.strip() for text in pages)
        return ContractDocument(
            source_format="pdf",
            text_es="\n".join(pages),
            page_count=pdf.page_count,
            needs_ocr=needs_ocr,
        )
    finally:
        pdf.close()


def _extract_docx(data: bytes) -> ContractDocument:
    try:
        document = docx.Document(BytesIO(data))
    except Exception as exc:
        raise DocumentUnreadable("DOCX cannot be opened") from exc
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    if not text.strip():
        raise DocumentUnreadable("DOCX without text")
    return ContractDocument(
        source_format="docx", text_es=text, page_count=1, needs_ocr=False
    )


def extract(data: bytes) -> ContractDocument:
    if data.startswith(b"%PDF"):
        return _extract_pdf(data)
    if data.startswith(b"PK\x03\x04"):
        return _extract_docx(data)
    raise UnsupportedFormat("PDF or DOCX expected")
